"""Inventory import (CSV/XLSX → dry run → apply) and export with the same columns (plan §6.3).

Semantics of a mapped cell:
* nullable fields (price, plan, notes, room counts…): an empty cell means "no value" (NULL),
  so a missing price survives an export → import round trip as NULL, never as 0;
* non-nullable fields (availability, currency, price_on_request): an empty cell means
  "keep the current value";
* an unmapped column never changes anything.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from unitforma.audit.service import write_audit
from unitforma.core.errors import ProblemError
from unitforma.core.ids import new_id
from unitforma.sales import rules
from unitforma.sales import service as sales_service
from unitforma.storage.service import get_storage, object_key, safe_filename
from unitforma.tenants.models import ImportBatch, Project, Unit, UnitSales

MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000
Actor = tuple[uuid.UUID, str]

# import targets in column order; label is what the export writes as the header
TARGETS: dict[str, str] = {
    "unit_id": "unit_id",
    "code": "code",
    "availability": "availability",
    "availability_reason": "availability_reason",
    "list_price": "list_price",
    "currency": "currency",
    "price_on_request": "price_on_request",
    "price_area_type": "price_area_type",
    "payment_plan": "payment_plan",
    "promotions": "promotions",
    "rooms": "rooms",
    "bedrooms": "bedrooms",
    "bathrooms": "bathrooms",
    "orientation": "orientation",
    "entrance": "entrance",
    "public_description": "public_description",
    "internal_notes": "internal_notes",
}
SYNONYMS: dict[str, tuple[str, ...]] = {
    "unit_id": ("id", "unit_uuid"),
    "code": ("unit", "unit_code", "apartment", "apt", "ბინა", "ბინის_კოდი"),
    "availability": ("status", "sales_status", "სტატუსი"),
    "list_price": ("price", "listprice", "ფასი"),
    "currency": ("ccy", "ვალუტა"),
    "price_on_request": ("on_request", "por"),
    "payment_plan": ("plan", "payment"),
    "promotions": ("promotion", "offers"),
    "public_description": ("description", "აღწერა"),
    "internal_notes": ("notes", "internal_note"),
    "entrance": ("staircase", "stair", "podezd", "სადარბაზო"),
}
NON_NULLABLE = ("availability", "currency", "price_on_request")
READ_ONLY_EXPORT = (
    "building",
    "floor",
    "unit_type",
    "verification_status",
    "publication_status",
    "price_per_m2",
)


def norm_header(h: Any) -> str:
    return re.sub(r"[\s\-/.]+", "_", str(h or "").strip().lower()).strip("_")


def auto_mapping(headers: list[str], saved: dict[str, str] | None = None) -> dict[str, str]:
    """target → header. Saved project mapping first, then exact names, then synonyms."""
    out: dict[str, str] = {}
    present = {h for h in headers if h}
    for target, header in (saved or {}).items():
        if target in TARGETS and header in present:
            out[target] = header
    by_norm = {norm_header(h): h for h in headers if h}
    for target in TARGETS:
        if target in out:
            continue
        for cand in (target, *SYNONYMS.get(target, ())):
            if cand in by_norm and by_norm[cand] not in out.values():
                out[target] = by_norm[cand]
                break
    return out


# --- file parsing ---------------------------------------------------------------------------
def detect_format(file_name: str, data: bytes) -> str:
    if data[:2] == b"PK" or file_name.lower().endswith(".xlsx"):
        return "xlsx"
    return "csv"


def parse_file(data: bytes, fmt: str) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    """Returns headers and (sheet row number, cells) for each non-empty data row."""
    if fmt == "xlsx":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ProblemError(422, "Could not read the XLSX file", str(e)[:300]) from e
        ws = wb.worksheets[0]
        raw = [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
        wb.close()
    else:
        text = None
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ProblemError(422, "Could not decode the CSV file")
        sample = text[:20000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [(i, r) for i, r in enumerate(csv.reader(io.StringIO(text), dialect), start=1)]
    rows = [(i, r) for i, r in raw if any(not rules.is_blank(c) for c in r)]
    if not rows:
        raise ProblemError(422, "The file has no rows")
    header_no, header = rows[0]
    headers = [str(h).strip() if h is not None else "" for h in header]
    if len(set(h for h in headers if h)) != len([h for h in headers if h]):
        raise ProblemError(422, "The header row has duplicate column names")
    body = rows[1:]
    if len(body) > MAX_ROWS:
        raise ProblemError(413, f"At most {MAX_ROWS} rows per import")
    return headers, body


# --- batches ----------------------------------------------------------------------------------
async def upload(
    db: AsyncSession, project: Project, file_name: str, data: bytes, actor: Actor
) -> ImportBatch:
    if not data:
        raise ProblemError(422, "Empty file")
    if len(data) > MAX_IMPORT_BYTES:
        raise ProblemError(413, "Import files are limited to 20 MB")
    fmt = detect_format(file_name, data)
    headers, body = parse_file(data, fmt)
    bid = new_id()
    key = object_key(
        project.tenant_id, project.id, "imports", f"{bid.hex[:12]}-{safe_filename(file_name)}"
    )
    get_storage().put_bytes(
        key,
        data,
        "text/csv"
        if fmt == "csv"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    saved = (project.sales_settings or {}).get("import_mapping") or {}
    batch = ImportBatch(
        id=bid,
        tenant_id=project.tenant_id,
        project_id=project.id,
        file_name=safe_filename(file_name),
        object_key=key,
        file_format=fmt,
        status="uploaded",
        headers=headers,
        mapping=auto_mapping(headers, saved),
        options=dict((project.sales_settings or {}).get("import_options") or {}),
        row_count=len(body),
        created_by=actor[0],
    )
    db.add(batch)
    await db.flush()
    return batch


async def get_batch(db: AsyncSession, tenant_id: uuid.UUID, batch_id: uuid.UUID) -> ImportBatch:
    b = (
        await db.execute(
            select(ImportBatch).where(
                ImportBatch.tenant_id == tenant_id, ImportBatch.id == batch_id
            )
        )
    ).scalar_one_or_none()
    if b is None:
        raise ProblemError(404, "Import not found")
    return b


def _check_mapping(headers: list[str], mapping: dict[str, str]) -> None:
    unknown = [t for t in mapping if t not in TARGETS]
    if unknown:
        raise ProblemError(422, f"Unknown import field(s): {', '.join(unknown)}")
    missing = [h for h in mapping.values() if h not in headers]
    if missing:
        raise ProblemError(422, f"Column(s) not in the file: {', '.join(missing)}")
    if len(set(mapping.values())) != len(mapping):
        raise ProblemError(422, "A column is mapped to more than one field")
    if "unit_id" not in mapping and "code" not in mapping:
        raise ProblemError(422, "Map the unit_id or the code column to identify units")


def _options(options: dict[str, Any]) -> tuple[rules.NumberFormat, bool]:
    try:
        fmt = rules.NumberFormat(
            decimal_separator=options.get("decimal_separator", "."),
            thousands_separator=options.get("thousands_separator", ""),
        )
    except ValueError as e:
        raise ProblemError(422, "Invalid number format", str(e)) from e
    return fmt, bool(options.get("confirm_zero_price", False))


async def _compute(
    db: AsyncSession,
    project: Project,
    headers: list[str],
    body: list[tuple[int, list[Any]]],
    mapping: dict[str, str],
    options: dict[str, Any],
) -> tuple[dict[str, Any], list[tuple[int, Unit, UnitSales | None, list[rules.FieldDiff]]]]:
    fmt, confirm_zero = _options(options)
    lookups = await sales_service.load_lookups(db, project)
    units = list((await db.execute(select(Unit).where(Unit.project_id == project.id))).scalars())
    by_id = {u.id: u for u in units}
    by_code = {u.code.lower(): u for u in units if u.lifecycle_status != "removed"}
    sales_rows = {
        s.unit_id: s
        for s in (
            await db.execute(select(UnitSales).where(UnitSales.project_id == project.id))
        ).scalars()
    }
    col = {target: headers.index(h) for target, h in mapping.items()}
    report_rows: list[dict[str, Any]] = []
    plan: list[tuple[int, Unit, UnitSales | None, list[rules.FieldDiff]]] = []
    seen: dict[uuid.UUID, int] = {}
    unknown_codes: list[str] = []
    by_field: dict[str, int] = {}

    def cell(cells: list[Any], target: str) -> Any:
        i = col.get(target)
        return cells[i] if i is not None and i < len(cells) else None

    for row_no, cells in body:
        errors: list[str] = []
        warnings: list[str] = []
        unit: Unit | None = None
        raw_id, raw_code = cell(cells, "unit_id"), cell(cells, "code")
        if not rules.is_blank(raw_id):
            try:
                unit = by_id.get(uuid.UUID(str(raw_id).strip()))
            except ValueError:
                errors.append(f"unit_id '{raw_id}' is not a valid id")
            else:
                if unit is None:
                    errors.append(f"unit_id {raw_id} does not belong to this project")
                elif unit.lifecycle_status == "removed":
                    errors.append(f"unit {unit.code} is removed")
                    unit = None
            if unit is not None and not rules.is_blank(raw_code):
                if str(raw_code).strip().lower() != unit.code.lower():
                    warnings.append(
                        f"code '{raw_code}' differs from the unit's current code {unit.code}; "
                        "codes are changed in the mapping workspace, not by import"
                    )
        elif not rules.is_blank(raw_code):
            unit = by_code.get(str(raw_code).strip().lower())
            if unit is None:
                errors.append(f"unknown code '{raw_code}'")
                unknown_codes.append(str(raw_code).strip())
        else:
            errors.append("row has neither unit_id nor code")
        if unit is not None and unit.id in seen:
            errors.append(f"unit {unit.code} already appears in row {seen[unit.id]}")
            unit = None
        if unit is not None:
            seen[unit.id] = row_no

        changes: dict[str, Any] = {}
        explicit_por: bool | None = None
        for target in mapping:
            if target in ("unit_id", "code"):
                continue
            v = cell(cells, target)
            try:
                if target in NON_NULLABLE and rules.is_blank(v):
                    continue
                if target == "availability":
                    changes["availability"] = rules.parse_availability(v)
                elif target == "availability_reason":
                    changes["availability_reason"] = rules.parse_text(v, 2000)
                elif target == "list_price":
                    changes["list_price"] = rules.parse_price(v, fmt)
                elif target == "currency":
                    changes["currency"] = rules.parse_currency(v)
                elif target == "price_on_request":
                    explicit_por = rules.parse_bool(v)
                    changes["price_on_request"] = explicit_por
                elif target == "price_area_type":
                    changes["price_area_type"] = rules.parse_text(v, 40)
                elif target == "payment_plan":
                    name = rules.parse_text(v, 120)
                    if name is None:
                        changes["payment_plan_id"] = None
                    elif name.lower() in lookups.plan_by_name:
                        changes["payment_plan_id"] = lookups.plan_by_name[name.lower()]
                    else:
                        raise rules.CellError(f"unknown payment plan '{name}'")
                elif target == "promotions":
                    text = rules.parse_text(v, 2000)
                    ids = []
                    for name in [n.strip() for n in (text or "").split(";") if n.strip()]:
                        if name.lower() not in lookups.promotion_by_name:
                            raise rules.CellError(f"unknown promotion '{name}'")
                        ids.append(lookups.promotion_by_name[name.lower()])
                    changes["promotion_ids"] = ids
                elif target in ("rooms", "bedrooms", "bathrooms"):
                    changes[target] = rules.parse_int(v, fmt)
                elif target in ("orientation", "entrance"):
                    changes[target] = rules.parse_text(v, 20)
                elif target in ("public_description", "internal_notes"):
                    changes[target] = rules.parse_text(v, 4000)
            except rules.CellError as e:
                errors.append(f"{mapping[target]}: {e}")

        if changes.get("list_price") is not None and "currency" not in changes:
            errors.append("a currency is required when a price is given")

        diffs: list[rules.FieldDiff] = []
        version = 0
        if unit is not None:
            sales = sales_rows.get(unit.id)
            cur = rules.current_values(project, unit, sales)
            version = cur.version
            diffs = rules.diff(cur, changes)
            allow_zero = confirm_zero and explicit_por is False
            errors += rules.validate(
                diffs,
                cur,
                lookups,
                allow_zero_price=allow_zero,
                reason=changes.get("availability_reason"),
            )
            if not errors and diffs:
                plan.append((row_no, unit, sales, diffs))
                for d in diffs:
                    by_field[d.field] = by_field.get(d.field, 0) + 1
        report_rows.append(
            {
                "row": row_no,
                "unit_id": str(unit.id) if unit else None,
                "code": unit.code if unit else (str(raw_code).strip() if raw_code else None),
                "status": "error" if errors else ("changed" if diffs else "unchanged"),
                "changes": [d.as_json() for d in diffs],
                "errors": errors,
                "warnings": warnings,
                "version": version,
            }
        )
    summary = {
        "rows": len(report_rows),
        "changed": sum(1 for r in report_rows if r["status"] == "changed"),
        "unchanged": sum(1 for r in report_rows if r["status"] == "unchanged"),
        "errors": sum(1 for r in report_rows if r["status"] == "error"),
        "warnings": sum(1 for r in report_rows if r["warnings"]),
        "unknown_codes": unknown_codes,
        "changes_by_field": by_field,
        "total_changes": sum(by_field.values()),
    }
    return {"summary": summary, "rows": report_rows}, plan


async def dry_run(
    db: AsyncSession,
    project: Project,
    batch: ImportBatch,
    mapping: dict[str, str],
    options: dict[str, Any],
    actor: Actor,
) -> ImportBatch:
    if batch.status in ("applied", "discarded"):
        raise ProblemError(409, f"Import is already {batch.status}")
    _check_mapping(batch.headers, mapping)
    headers, body = parse_file(get_storage().get_bytes(batch.object_key), batch.file_format)
    report, _plan = await _compute(db, project, headers, body, mapping, options)
    batch.mapping = mapping
    batch.options = options
    batch.report = report
    batch.status = "dry_run"
    batch.row_count = report["summary"]["rows"]
    batch.change_count = report["summary"]["total_changes"]
    batch.error_count = report["summary"]["errors"]
    # the mapping and number format are remembered per project (plan §6.3)
    settings = dict(project.sales_settings or {})
    settings["import_mapping"] = mapping
    settings["import_options"] = {
        k: options[k] for k in ("decimal_separator", "thousands_separator") if k in options
    }
    project.sales_settings = settings
    return batch


async def apply(
    db: AsyncSession, project: Project, batch: ImportBatch, skip_invalid_rows: bool, actor: Actor
) -> ImportBatch:
    """Apply exactly what the dry run showed, in one transaction (all-or-nothing)."""
    if batch.status != "dry_run":
        raise ProblemError(409, "Run a dry run before applying", f"import is {batch.status}")
    headers, body = parse_file(get_storage().get_bytes(batch.object_key), batch.file_format)
    report, plan = await _compute(db, project, headers, body, batch.mapping, batch.options)
    seen_versions = {
        r["unit_id"]: r["version"] for r in batch.report.get("rows", []) if r["unit_id"]
    }
    now_versions = {r["unit_id"]: r["version"] for r in report["rows"] if r["unit_id"]}
    stale = [u for u, v in now_versions.items() if seen_versions.get(u) != v]
    if stale or report["summary"] != batch.report.get("summary"):
        raise ProblemError(
            409,
            "Inventory changed since the dry run",
            "Run the dry run again to see the current differences before applying",
            extra={"stale_units": stale[:100]},
        )
    if report["summary"]["errors"] and not skip_invalid_rows:
        raise ProblemError(
            422,
            "The import has errors",
            f"{report['summary']['errors']} row(s) have errors: fix the file, or apply with "
            "skip_invalid_rows to import only the valid rows",
        )
    changes = 0
    for _row_no, unit, sales, diffs in plan:
        locked = await sales_service.get_sales(db, unit.id, lock=True)
        await sales_service.apply_changes(
            db,
            project=project,
            unit=unit,
            sales=locked if locked is not None else sales,
            diffs=diffs,
            actor=actor,
            reason=f"import {batch.file_name}",
            import_batch_id=batch.id,
        )
        changes += len(diffs)
    batch.status = "applied"
    batch.applied_by = actor[0]
    batch.applied_at = datetime.now(UTC)
    batch.change_count = changes
    await write_audit(
        db,
        tenant_id=project.tenant_id,
        actor_user_id=actor[0],
        actor_role=actor[1],
        entity_type="import_batch",
        entity_id=batch.id,
        action="apply",
        after={
            "file_name": batch.file_name,
            "units": len(plan),
            "changes": changes,
            "skipped_rows": report["summary"]["errors"],
        },
    )
    return batch


# --- export -----------------------------------------------------------------------------------
def export_columns(project: Project, column_set: str) -> list[str]:
    cols = [
        "unit_id",
        "code",
        "building",
        "floor",
        "unit_type",
        "verification_status",
        "publication_status",
        "availability",
        "availability_reason",
        "list_price",
        "currency",
        "price_on_request",
        "price_area_type",
        "price_per_m2",
        "payment_plan",
        "promotions",
        "rooms",
        "bedrooms",
        "bathrooms",
        "orientation",
        "entrance",
        "public_description",
        "internal_notes",
    ]
    if column_set == "public":
        cols = [c for c in cols if c not in rules.INTERNAL_ONLY]
    keys = [a.get("key") for a in (project.area_definitions or []) if isinstance(a, dict)]
    return cols + [f"area_{k}" for k in keys]


def _export_value(row: dict[str, Any], col: str) -> Any:
    if col.startswith("area_"):
        return row["areas"].get(col[5:])
    return {
        "unit_id": str(row["unit_id"]),
        "building": row["building_code"],
        "floor": row["floor_level"],
        "price_area_type": row["price_area_type_override"],
        "payment_plan": row["payment_plan_name"],
        "promotions": "; ".join(row["promotion_names"]) or None,
    }.get(col, row.get(col))


async def export(
    db: AsyncSession, project: Project, file_format: str, column_set: str
) -> tuple[bytes, str, str]:
    rows = await sales_service.inventory(db, project)
    cols = export_columns(project, column_set)
    stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M")
    name = f"{project.slug}-inventory-{column_set}-{stamp}.{file_format}"
    if file_format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        w.writerow(cols)
        for r in rows:
            out = []
            for c in cols:
                v = _export_value(r, c)
                if isinstance(v, bool):
                    v = "true" if v else "false"
                out.append("" if v is None else rules.to_json(v))
            w.writerow(out)
        return buf.getvalue().encode("utf-8-sig"), name, "text/csv; charset=utf-8"
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    price_cols = {i for i, c in enumerate(cols, start=1) if c in ("list_price", "price_per_m2")}
    area_cols = {i for i, c in enumerate(cols, start=1) if c.startswith("area_")}
    for r in rows:
        ws.append([_export_value(r, c) for c in cols])
        n = ws.max_row
        for i in price_cols:
            ws.cell(n, i).number_format = "#,##0.00"
        for i in area_cols:
            ws.cell(n, i).number_format = "0.00"
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    return (
        buf.getvalue(),
        name,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
